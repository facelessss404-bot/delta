import openpyxl
import json
import re

# Exact 19 daily time periods from 05:00 to 23:00
TIME_SLOT_MAP = {
    1: ('05:00', '05:30', '05:00 - 05:30'),
    2: ('05:30', '07:30', '05:30 - 07:30'),
    3: ('07:30', '08:00', '07:30 - 08:00'),
    4: ('08:00', '08:30', '08:00 - 08:30'),
    5: ('08:30', '09:00', '08:30 - 09:00'),
    6: ('09:00', '10:00', '09:00 - 10:00'),
    7: ('10:00', '11:30', '10:00 - 11:30'),
    8: ('11:30', '11:45', '11:30 - 11:45'),
    9: ('11:45', '13:00', '11:45 - 13:00'),
    10: ('13:00', '13:30', '13:00 - 13:30'),
    11: ('13:30', '14:30', '13:30 - 14:30'),
    12: ('14:30', '15:45', '14:30 - 15:45'),
    13: ('15:45', '16:00', '15:45 - 16:00'),
    14: ('16:00', '17:00', '16:00 - 17:00'),
    15: ('17:00', '18:00', '17:00 - 18:00'),
    16: ('18:00', '20:30', '18:00 - 20:30'),
    17: ('20:30', '21:00', '20:30 - 21:00'),
    18: ('21:00', '22:30', '21:00 - 22:30'),
    19: ('22:30', '23:00', '22:30 - 23:00')
}

def parse_slot_text(val):
    val = (val or '').strip()
    if not val:
        return ('Free / Self-Study', '', 'study')
    val_upper = val.upper()
    if 'WAKE' in val_upper: return ('Wake Up', '', 'routine')
    if 'FRESH' in val_upper: return ('Fresh Up', '', 'routine')
    if 'FALL' in val_upper: return ('Fall In / Roll Call', '', 'routine')
    if 'BREAKFAST' in val_upper: return ('Breakfast', '', 'routine')
    if 'DINNER' in val_upper: return ('Dinner', '', 'routine')
    if 'LIGHT' in val_upper: return ('Lights Off', '', 'routine')
    if val_upper in ['B', 'R', 'E', 'A', 'K']: return ('Break / Refreshments', '', 'routine')
    if val_upper in ['L', 'U', 'N', 'C', 'H']: return ('Lunch Break', '', 'routine')
    if val_upper in ['G', 'A', 'M', 'E', 'S']: return ('Games & Sports', 'Devika / Rithika / Prabha / Abhi', 'games')
    if 'PHYSICAL' in val_upper or 'PHSICAL' in val_upper: return ('Physical Training (PT)', 'Training Staff', 'physical')
    if 'STUDY' in val_upper: return ('Study Hours', 'Supervised Study', 'study')
    if 'CURRENT' in val_upper or 'CUREENT' in val_upper: return ('Current Affairs & GK', '', 'academic')
    if 'LEARNING' in val_upper: return ('Learning Videos / Defence Prep', '', 'academic')
    if 'SELF' in val_upper: return ('Self Preparation & Revision', '', 'study')
    
    instructor = ''
    activity = val.replace('\n', ' ')
    m = re.search(r'\((.*?)\)', activity)
    if m:
        instructor = m.group(1).strip()
        activity = re.sub(r'\(.*?\)', '', activity).strip()
    activity = re.sub(r'\s+', ' ', activity).strip()
    return (activity, instructor, 'academic')

output_slots = []

# ==========================================
# 1. PARSE JLP SCHEDULE (Classes 8–12)
# ==========================================
wb_jlp = openpyxl.load_workbook('JLP SCHEDULE.xlsx', data_only=True)

# 1.1 Sheet1: Grade 8th - 10th
ws1 = wb_jlp['Sheet1']
for r in range(4, 10):
    day = ws1.cell(r, 1).value
    if not day: continue
    day_name = day.strip().capitalize()
    for col_idx in range(2, 21):
        slot_order = col_idx - 1
        val = ws1.cell(r, col_idx).value
        t1, t2, disp = TIME_SLOT_MAP[slot_order]
        act, inst, cat = parse_slot_text(val)
        output_slots.append({
            'schedule_key': 'jlp_8_10',
            'schedule_title': 'JLP: Classes 8th–10th',
            'programme': 'JLP',
            'day_of_week': day_name,
            'slot_order': slot_order,
            'start_time': t1,
            'end_time': t2,
            'time_display': disp,
            'activity': act,
            'instructor': inst,
            'category': cat,
            'raw_text': val or '',
            'trackable': cat in ['academic', 'physical', 'study', 'games']
        })

# 1.2 Sheet2: Grade 11th & 12th
ws2 = wb_jlp['Sheet2']
for r in range(5, 11):
    day = ws2.cell(r, 2).value
    if not day: continue
    day_name = day.strip().capitalize()
    for col_idx in range(3, 22):
        slot_order = col_idx - 2
        val = ws2.cell(r, col_idx).value
        t1, t2, disp = TIME_SLOT_MAP[slot_order]
        act, inst, cat = parse_slot_text(val)
        output_slots.append({
            'schedule_key': 'jlp_11_12',
            'schedule_title': 'JLP: Classes 11th–12th',
            'programme': 'JLP',
            'day_of_week': day_name,
            'slot_order': slot_order,
            'start_time': t1,
            'end_time': t2,
            'time_display': disp,
            'activity': act,
            'instructor': inst,
            'category': cat,
            'raw_text': val or '',
            'trackable': cat in ['academic', 'physical', 'study', 'games']
        })

# ==========================================
# 2. PARSE SLP SCHEDULE (Graduates)
# ==========================================
wb_slp = openpyxl.load_workbook("SLP's SCHEDULE.xlsx", data_only=True)
ws_slp = wb_slp['Sheet1']
for r in range(4, 10):
    day = ws_slp.cell(r, 1).value
    if not day: continue
    day_name = day.strip().capitalize()
    for col_idx in range(2, 21):
        slot_order = col_idx - 1
        val = ws_slp.cell(r, col_idx).value
        t1, t2, disp = TIME_SLOT_MAP[slot_order]
        act, inst, cat = parse_slot_text(val)
        output_slots.append({
            'schedule_key': 'slp',
            'schedule_title': 'SLP: Graduates',
            'programme': 'SLP',
            'day_of_week': day_name,
            'slot_order': slot_order,
            'start_time': t1,
            'end_time': t2,
            'time_display': disp,
            'activity': act,
            'instructor': inst,
            'category': cat,
            'raw_text': val or '',
            'trackable': cat in ['academic', 'physical', 'study', 'games']
        })

print(f'Total slots extracted: {len(output_slots)}')
by_key = {}
for s in output_slots:
    by_key[s['schedule_key']] = by_key.get(s['schedule_key'], 0) + 1
print('Breakdown by schedule key:', by_key)

with open('timetable_data.json', 'w', encoding='utf-8') as f:
    json.dump(output_slots, f, indent=2, ensure_ascii=False)

print('Saved timetable_data.json successfully!')
